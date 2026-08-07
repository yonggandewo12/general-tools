"""图表创建（line/bar/pie/scatter/area）。

依赖 openpyxl 原生图表能力，exceljs 无对应功能，故保留 Python 实现。
"""

from __future__ import annotations

import logging
import re
from typing import Any

from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.utils import column_index_from_string

from ._utils import edit_workbook, parse_cell_range, require_sheet
from .exceptions import ChartError, ValidationError

logger = logging.getLogger(__name__)

_CHART_CLASSES = {
    "line": LineChart,
    "bar": BarChart,
    "pie": PieChart,
    "scatter": ScatterChart,
    "area": AreaChart,
}

_TARGET_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")


def _parse_target_cell(target_cell: str) -> tuple[int, int]:
    """解析图表锚点单元格为 (col_index_0based, row_index_0based)。"""
    m = _TARGET_RE.match(target_cell.upper())
    if not m:
        raise ValidationError(f"Invalid target cell format: {target_cell}")
    col = column_index_from_string(m.group(1)) - 1
    row = int(m.group(2)) - 1
    return col, row


def create_chart_in_sheet(
    filepath: str,
    sheet_name: str,
    data_range: str,
    chart_type: str,
    target_cell: str,
    title: str = "",
    x_axis: str = "",
    y_axis: str = "",
    style: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """在工作表中创建图表。"""
    if style is None:
        style = {"show_data_labels": True}
    else:
        style.setdefault("show_data_labels", True)
    try:
        chart_type_lower = chart_type.lower()
        chart_cls = _CHART_CLASSES.get(chart_type_lower)
        if chart_cls is None:
            raise ValidationError(
                f"Unsupported chart type: {chart_type}. Supported: {', '.join(_CHART_CLASSES)}"
            )

        # 解析数据范围（可带 sheet! 前缀）
        if "!" in data_range:
            range_sheet, cell_range = data_range.split("!", 1)
        else:
            range_sheet, cell_range = sheet_name, data_range
        if ":" not in cell_range:
            raise ValidationError(f"Invalid data range format: {data_range}")
        start_cell, end_cell = cell_range.split(":", 1)
        sr, sc, er, ec = parse_cell_range(start_cell, end_cell)

        with edit_workbook(filepath) as wb:
            if range_sheet not in wb.sheetnames:
                raise ValidationError(f"Sheet {range_sheet!r} referenced in data range not found")
            ws = wb[range_sheet]
            if sheet_name not in wb.sheetnames:
                raise ValidationError(f"Sheet {sheet_name!r} not found")
            target_ws = wb[sheet_name]

            chart = chart_cls()
            chart.title = title
            if hasattr(chart, "x_axis"):
                chart.x_axis.title = x_axis
            if hasattr(chart, "y_axis"):
                chart.y_axis.title = y_axis

            try:
                if chart_type_lower == "scatter":
                    for col in range(sc + 1, ec + 1):
                        x_values = Reference(ws, min_row=sr + 1, max_row=er, min_col=sc)
                        y_values = Reference(ws, min_row=sr + 1, max_row=er, min_col=col)
                        chart.series.append(Series(y_values, x_values, title_from_data=True))
                else:
                    data = Reference(ws, min_row=sr, max_row=er, min_col=sc + 1, max_col=ec)
                    cats = Reference(ws, min_row=sr + 1, max_row=er, min_col=sc)
                    chart.add_data(data, titles_from_data=True)
                    chart.set_categories(cats)
            except Exception as e:
                raise ChartError(f"Failed to create chart data references: {e}") from e

            _apply_chart_style(chart, style)
            chart.width = 15
            chart.height = 7.5

            try:
                _parse_target_cell(target_cell)
                target_ws.add_chart(chart, target_cell)
            except Exception as e:
                raise ChartError(f"Failed to anchor chart: {e}") from e

        return {
            "message": f"{chart_type.capitalize()} chart created successfully",
            "details": {"type": chart_type, "location": target_cell, "data_range": data_range},
        }
    except (ValidationError, ChartError):
        raise
    except Exception as e:
        logger.error("Unexpected error creating chart: %s", e)
        raise ChartError(str(e)) from e


def _apply_chart_style(chart, style: dict[str, Any]) -> None:
    try:
        if style.get("show_legend", True):
            chart.legend = Legend()
            chart.legend.position = style.get("legend_position", "r")
        else:
            chart.legend = None

        if style.get("show_data_labels", False):
            dl = DataLabelList()
            opts = style.get("data_label_options", {}) if isinstance(style.get("data_label_options"), dict) else {}

            def opt(name: str, default: bool) -> bool:
                return bool(opts.get(name, default))

            dl.showVal = opt("show_val", True)
            dl.showCatName = opt("show_cat_name", False)
            dl.showSerName = opt("show_ser_name", False)
            dl.showLegendKey = opt("show_legend_key", False)
            dl.showPercent = opt("show_percent", False)
            dl.showBubbleSize = opt("show_bubble_size", False)
            chart.dataLabels = dl

        if style.get("grid_lines", False):
            if hasattr(chart, "x_axis"):
                chart.x_axis.majorGridlines = ChartLines()
            if hasattr(chart, "y_axis"):
                chart.y_axis.majorGridlines = ChartLines()
    except Exception as e:
        raise ChartError(f"Failed to apply chart style: {e}") from e
